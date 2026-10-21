#!/usr/bin/env python3
"""
Mirror data/jobs.xlsx into a Google Sheet so Michelle can view/edit it anywhere
(including her phone). Runs right after poller.py each cycle.

Design: APPEND-ONLY. It reads the URLs already in the Sheet and only adds rows
whose URL isn't there yet. That means any edits you make in the Sheet — the
Status column especially — are never overwritten. The local xlsx stays the full
record; the Sheet is your live, editable view.

Setup (one time): create a Google service account + JSON key, drop the key at
gcp-service-account.json in this folder, create a blank Google Sheet, share it
(Editor) with the service-account email, and put its ID in sheet_config.json:
    { "sheet_id": "<the long id from the Sheet URL>" }
If either file is missing, this script no-ops quietly so it never breaks the poller.
"""
import json
import os
import sys
import socket

# Never let a Google API call block forever. gspread -> requests -> urllib3 sets
# no timeout by default, so a half-open socket (laptop sleep, WiFi change, VPN)
# hangs the read indefinitely. That stalls run.sh, and launchd will not start a
# new cycle while the old one is alive -- one hang silently freezes the tracker.
socket.setdefaulttimeout(60)

HERE = os.path.dirname(os.path.abspath(__file__))
KEY_PATH = os.path.join(HERE, "gcp-service-account.json")
CONFIG_PATH = os.path.join(HERE, "sheet_config.json")
XLSX_PATH = os.path.join(HERE, "data", "jobs.xlsx")
SCOPES = ["https://www.googleapis.com/auth/spreadsheets"]
HEADER = ["FirstSeen (ET)", "Company", "Title", "Category", "Level",
          "Location", "Comp (if listed)", "Posted", "URL", "Status"]
URL_COL = 9  # 1-indexed column holding the job URL (our dedupe key)


def log(msg):
    print(f"[sheets_sync] {msg}")


def main():
    if not os.path.exists(KEY_PATH) or not os.path.exists(CONFIG_PATH):
        log("no credentials/config yet — skipping (set up gcp-service-account.json "
            "+ sheet_config.json to enable).")
        return
    if not os.path.exists(XLSX_PATH):
        log("no jobs.xlsx yet — nothing to sync.")
        return

    import gspread
    from google.oauth2.service_account import Credentials
    from openpyxl import load_workbook

    sheet_id = json.load(open(CONFIG_PATH))["sheet_id"]
    creds = Credentials.from_service_account_file(KEY_PATH, scopes=SCOPES)
    gc = gspread.authorize(creds)
    sh = gc.open_by_key(sheet_id)

    try:
        ws = sh.worksheet("Jobs")
    except gspread.WorksheetNotFound:
        ws = sh.add_worksheet("Jobs", rows=1000, cols=len(HEADER))
        ws.update([HEADER], "A1")
        ws.format("A1:J1", {"textFormat": {"bold": True}})
        ws.freeze(rows=1)

    # URLs already in the Sheet (skip header)
    existing = set()
    col = ws.col_values(URL_COL)[1:]
    for u in col:
        if u:
            existing.add(u.strip())

    # rows from the local xlsx (source of record)
    wb = load_workbook(XLSX_PATH)
    xws = wb["Jobs"]
    to_append = []
    for row in xws.iter_rows(min_row=2, values_only=True):
        row = list(row) + [""] * (len(HEADER) - len(row))
        url = (row[URL_COL - 1] or "").strip()
        if url and url not in existing:
            to_append.append([("" if v is None else v) for v in row[:len(HEADER)]])
            existing.add(url)

    if to_append:
        # newest-first: insert just under the header
        ws.insert_rows(to_append, row=2, value_input_option="USER_ENTERED")
        log(f"appended {len(to_append)} new rows to the Sheet.")
    else:
        log("Sheet already up to date — 0 new rows.")

    # Heartbeat so you can always see it's alive, even on a 0-new cycle.
    from datetime import datetime
    from zoneinfo import ZoneInfo
    if ws.col_count < 13:
        ws.add_cols(13 - ws.col_count)
    stamp = datetime.now(ZoneInfo("America/New_York")).strftime("%Y-%m-%d %H:%M ET")
    total = len(ws.col_values(URL_COL)) - 1
    ws.update([["Last checked:", stamp],
               ["New this cycle:", len(to_append)],
               ["Total roles:", total]], "L1", value_input_option="USER_ENTERED")
    ws.format("L1:L3", {"textFormat": {"bold": True}})


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        log(f"ERROR: {e}")
        sys.exit(0)  # never break the poller pipeline
