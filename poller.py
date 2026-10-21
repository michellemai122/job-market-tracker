#!/usr/bin/env python3
"""
Job poller. Reads registry.json (from discover.py), fetches each company's live
ATS feed, keeps only new-grad / APM / AI-engineer / forward-deployed roles in the
US (or remote), dedupes against data/seen.json, and appends brand-new postings to
data/jobs.xlsx — the continuously-updated spreadsheet.

Run it every 15-30 min from cron/launchd. Each run only adds jobs it hasn't seen
before, so the sheet grows and never loses your notes in the Status column.
"""
import json
import os
import re
import ssl
import sys
from datetime import datetime
from zoneinfo import ZoneInfo
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError

try:
    import certifi
    SSL_CTX = ssl.create_default_context(cafile=certifi.where())
except ImportError:
    SSL_CTX = ssl.create_default_context()

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "data")
os.makedirs(DATA, exist_ok=True)
SEEN_PATH = os.path.join(DATA, "seen.json")
XLSX_PATH = os.path.join(DATA, "jobs.xlsx")
REGISTRY_PATH = os.path.join(HERE, "registry.json")
ET = ZoneInfo("America/New_York")
TIMEOUT = 15
UA = {"User-Agent": "Mozilla/5.0 (personal job tracker)"}

COLUMNS = ["FirstSeen (ET)", "Company", "Title", "Category", "Level",
           "Location", "Comp (if listed)", "Posted", "URL", "Status"]

# Email alert settings live in the gitignored sheet_config.json (see sheet_config.example.json).
# Sends through Mail.app (default account) — a missing or empty alert_email disables it.
try:
    with open(os.path.join(HERE, "sheet_config.json")) as _f:
        _CONFIG = json.load(_f)
except (OSError, ValueError):
    _CONFIG = {}
ALERT_EMAIL = _CONFIG.get("alert_email", "")
SHEET_URL = f"https://docs.google.com/spreadsheets/d/{_CONFIG.get('sheet_id', '')}/edit"

SEND_MAIL_APPLESCRIPT = '''
on run argv
    set toAddr to item 1 of argv
    set theSubject to item 2 of argv
    set theBody to item 3 of argv
    tell application "Mail"
        set msg to make new outgoing message with properties {subject:theSubject, content:theBody, visible:false}
        tell msg to make new to recipient at end of to recipients with properties {address:toAddr}
        send msg
    end tell
end run
'''


def send_email_alert(new_rows):
    """Email the newly found roles (with links) via Mail.app. APM roles first."""
    if not ALERT_EMAIL or not new_rows or sys.platform != "darwin":
        return
    import subprocess
    rows = sorted(new_rows, key=lambda x: (x["category"] != "APM", x["company"]))
    lines = []
    for j in rows:
        comp = f"  [{j['comp']}]" if j.get("comp") else ""
        loc = f" — {j['location']}" if j.get("location") else ""
        lines.append(f"• {j['category']}: {j['company']} — {j['title']}{loc}{comp}\n  {j['url']}")
    subject = f"Job Tracker: {len(rows)} new role{'s' if len(rows) > 1 else ''} — " \
              f"{rows[0]['company']}: {rows[0]['title'][:50]}"
    body = "\n\n".join(lines) + f"\n\nFull sheet: {SHEET_URL}\n"
    try:
        subprocess.run(["osascript", "-", ALERT_EMAIL, subject, body],
                       input=SEND_MAIL_APPLESCRIPT, text=True, timeout=60, check=False)
    except Exception:
        pass

# ---------------------------------------------------------------- classification
SENIOR = re.compile(r"\b(senior|sr\.?|staff|principal|\blead\b|director|"
                    r"vp|head of|architect|\bii\b|\biii\b|\biv\b|distinguished|"
                    r"experienced|expert|\d\+\s*years)\b", re.I)
# A role only counts as NEW GRAD if the title carries one of these signals.
NEWGRAD = re.compile(r"(new\s*grad|new\s*graduate|new\s*college\s*grad|"
                     r"university\s*grad(uate)?|university\s*hire|campus|"
                     r"early[\s-]*career|entry[\s-]*level|recent\s*grad(uate)?|"
                     r"class\s*of\s*20\d\d|rotational|rotation\s*program|"
                     r"graduate\s*(program|scheme|engineer|analyst|developer|"
                     r"software|rotational)|grad(uate)?\s*20\d\d|20(26|27)\s*grad|"
                     r"\b(20\s*)?27\s*(new\s*grad|grad|start)|\bnew\s*grad\b)", re.I)
INTERN = re.compile(r"\b(intern|internship|co-?op|summer\s*20\d\d)\b", re.I)

APM = re.compile(r"\b(a ?pm|associate\s+product\s+manager|product\s+manager[,\s-]*"
                 r"(new\s*grad|university|early|rotational|associate)|rotational\s+"
                 r"product|product\s+management\s+(new\s*grad|rotational|associate))\b", re.I)
FDE = re.compile(r"forward[\s-]*deployed", re.I)
AIENG = re.compile(r"\b(ai\s+engineer|a\.i\.\s+engineer|machine\s+learning\s+engineer|"
                   r"ml\s+engineer|applied\s+(ai|ml)|member\s+of\s+technical\s+staff|"
                   r"gen(erative)?\s*ai\s+engineer)\b", re.I)
SWE = re.compile(r"\b(software\s+engineer|software\s+developer|swe|quant\w*\s+(researcher|"
                 r"developer|trader|engineer|analyst)|product\s+manager)\b", re.I)


def classify(title):
    """Return (category, 'New Grad') only for genuine new-grad roles in Michelle's
    families; else None. Every non-APM role must carry an explicit new-grad signal."""
    t = title or ""
    if INTERN.search(t):
        return None  # internships are not new-grad full-time roles
    # APM / Associate PM / rotational PM: inherently early-career -> always new grad
    if APM.search(t) and not SENIOR.search(t):
        return ("APM", "New Grad")
    # everything else must explicitly signal new grad, and must not be senior
    if not NEWGRAD.search(t) or SENIOR.search(t):
        return None
    if FDE.search(t):
        return ("Forward-Deployed Eng", "New Grad")
    if AIENG.search(t):
        return ("AI/ML Engineer", "New Grad")
    if SWE.search(t):
        cat = "New-Grad SWE/Quant" if re.search(r"engineer|develop|swe|quant", t, re.I) else "New-Grad PM"
        return (cat, "New Grad")
    return None


def classify_role(title):
    """For sources already curated as new-grad (e.g. the Simplify list): match by
    role family only — no in-title new-grad signal required — but drop interns/senior."""
    t = title or ""
    if INTERN.search(t) or SENIOR.search(t):
        return None
    if APM.search(t):
        return ("APM", "New Grad")
    if FDE.search(t):
        return ("Forward-Deployed Eng", "New Grad")
    if AIENG.search(t):
        return ("AI/ML Engineer", "New Grad")
    if SWE.search(t):
        cat = "New-Grad SWE/Quant" if re.search(r"engineer|develop|swe|quant", t, re.I) else "New-Grad PM"
        return (cat, "New Grad")
    return None


# Concrete US signals (NOT bare "remote" — "Remote - India" must not read as US)
US_STRONG = re.compile(r"\b(united states|usa|u\.s\.a?\.?|\bus\b|new york|nyc|"
                       r"san francisco|seattle|austin|boston|chicago|los angeles|"
                       r"mountain view|menlo|sunnyvale|palo alto|bellevue|denver|"
                       r"atlanta|dallas|houston|washington|\bd\.?c\.?\b|new jersey|"
                       r"california|virginia|texas|florida|arizona|georgia|"
                       r"\bca\b|\bny\b|\btx\b|\bwa\b|\bma\b|\bil\b|\bva\b|\bnj\b)", re.I)
NON_US = re.compile(r"\b(london|dublin|ireland|toronto|vancouver|montreal|canada|"
                    r"bangalore|bengaluru|hyderabad|gurgaon|gurugram|pune|noida|"
                    r"chennai|mumbai|delhi|india|singapore|tokyo|japan|sydney|"
                    r"melbourne|australia|amsterdam|netherlands|berlin|munich|"
                    r"germany|paris|france|tel aviv|israel|warsaw|poland|"
                    r"hong kong|shanghai|beijing|shenzhen|china|seoul|korea|"
                    r"mexico|brazil|sao paulo|zurich|switzerland|stockholm|sweden|"
                    r"madrid|barcelona|spain|milan|italy|belgrade|serbia|cyprus|"
                    r"limassol|lisbon|portugal|dubai|uae|bulgaria|romania|"
                    r"united kingdom|\buk\b|emea|apac|latam)\b", re.I)


def location_ok(loc):
    loc = loc or ""
    if not loc.strip():
        return True  # unknown -> keep, don't miss
    if NON_US.search(loc):
        # foreign location named -> keep only if it ALSO names a concrete US place
        return bool(US_STRONG.search(loc))
    return True


# Require a real salary figure: comma-thousands ($204,000) or K-suffix ($204K).
# Bare "$204" is rejected so we don't surface truncated junk.
COMP = re.compile(r"\$\s?\d{2,3}(?:,\d{3}|\s?[kK])"
                  r"(?:\s?(?:[-–—]|to)\s?\$?\s?\d{2,3}(?:,\d{3}|\s?[kK])?)?")


def extract_comp(text):
    if not text:
        return ""
    m = COMP.search(text)
    return m.group(0).strip() if m else ""


# ---------------------------------------------------------------- fetchers
def fetch_json(url):
    try:
        with urlopen(Request(url, headers=UA), timeout=TIMEOUT, context=SSL_CTX) as r:
            return json.loads(r.read().decode("utf-8", "replace"))
    except Exception:
        return None


def norm(company, provider, slug, uid, title, location, url, posted, comp):
    return {"company": company, "provider": provider, "slug": slug,
            "jid": str(uid), "uid": f"{company}:{uid}", "title": title,
            "location": location or "", "url": url, "posted": posted or "", "comp": comp or ""}


def from_greenhouse(company, slug):
    d = fetch_json(f"https://boards-api.greenhouse.io/v1/boards/{slug}/jobs")
    out = []
    if isinstance(d, dict):
        for j in d.get("jobs", []):
            out.append(norm(company, "greenhouse", slug, j.get("id"), j.get("title", ""),
                            (j.get("location") or {}).get("name", ""),
                            j.get("absolute_url", ""), j.get("updated_at", ""), ""))
    return out


def from_lever(company, slug):
    d = fetch_json(f"https://api.lever.co/v0/postings/{slug}?mode=json")
    out = []
    if isinstance(d, list):
        for j in d:
            loc = (j.get("categories") or {}).get("location", "")
            posted = j.get("createdAt", "")
            if isinstance(posted, (int, float)):
                posted = datetime.fromtimestamp(posted / 1000, ET).strftime("%Y-%m-%d")
            comp = extract_comp(j.get("descriptionPlain", ""))
            out.append(norm(company, "lever", slug, j.get("id"), j.get("text", ""), loc,
                            j.get("hostedUrl", ""), posted, comp))
    return out


def from_ashby(company, slug):
    d = fetch_json(f"https://api.ashbyhq.com/posting-api/job-board/{slug}?includeCompensation=true")
    out = []
    if isinstance(d, dict):
        for j in d.get("jobs", []):
            comp = ""
            c = j.get("compensation") or {}
            summ = c.get("compensationTierSummary") or c.get("summaryComponents")
            if isinstance(summ, str):
                comp = extract_comp(summ)
            out.append(norm(company, "ashby", slug, j.get("id"), j.get("title", ""),
                            j.get("location", ""),
                            j.get("jobUrl") or j.get("applyUrl", ""),
                            j.get("publishedAt", ""), comp))
    return out


def from_smartrecruiters(company, slug):
    d = fetch_json(f"https://api.smartrecruiters.com/v1/companies/{slug}/postings?limit=100")
    out = []
    if isinstance(d, dict):
        for j in d.get("content", []):
            loc = j.get("location") or {}
            locstr = ", ".join(x for x in [loc.get("city"), loc.get("region"), loc.get("country")] if x)
            uid = j.get("id")
            url = f"https://jobs.smartrecruiters.com/{slug}/{uid}"
            out.append(norm(company, "smartrecruiters", slug, uid, j.get("name", ""),
                            locstr, url, j.get("releasedDate", ""), ""))
    return out


FETCHERS = {"greenhouse": from_greenhouse, "lever": from_lever,
            "ashby": from_ashby, "smartrecruiters": from_smartrecruiters}


def fetch_text(url):
    try:
        with urlopen(Request(url, headers=UA), timeout=TIMEOUT, context=SSL_CTX) as r:
            return r.read().decode("utf-8", "replace")
    except Exception:
        return ""


def enrich_comp(job):
    """For a matched job with no comp yet, open its detail page and pull a salary range.
    Only called on the handful of matched roles, so the extra fetches are cheap."""
    if job.get("comp"):
        return job["comp"]
    prov, slug, jid = job["provider"], job["slug"], job["jid"]
    text = ""
    if prov == "greenhouse":
        text = fetch_text(f"https://boards-api.greenhouse.io/v1/boards/{slug}/jobs/{jid}?questions=false")
    elif prov == "smartrecruiters":
        text = fetch_text(f"https://api.smartrecruiters.com/v1/companies/{slug}/postings/{jid}")
    elif prov == "lever":
        text = fetch_text(f"https://api.lever.co/v0/postings/{slug}/{jid}?mode=json")
    # ashby comp already handled in list fetch
    return extract_comp(text)


# ---------------------------------------------------------------- spreadsheet
def load_seen():
    if os.path.exists(SEEN_PATH):
        with open(SEEN_PATH) as f:
            return set(json.load(f))
    return set()


def save_seen(seen):
    with open(SEEN_PATH, "w") as f:
        json.dump(sorted(seen), f)


def ensure_workbook():
    if os.path.exists(XLSX_PATH):
        return load_workbook(XLSX_PATH)
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append(COLUMNS)
    header_fill = PatternFill("solid", fgColor="1F2937")
    for i, _ in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill
    ws.freeze_panes = "A2"
    widths = [18, 16, 42, 20, 10, 26, 20, 12, 60, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return wb


def main():
    if not os.path.exists(REGISTRY_PATH):
        sys.exit("registry.json not found — run discover.py first.")
    with open(REGISTRY_PATH) as f:
        registry = json.load(f)["found"]

    seen = load_seen()
    all_jobs = []
    with ThreadPoolExecutor(max_workers=12) as ex:
        futs = {ex.submit(FETCHERS[r["provider"]], r["company"], r["slug"]): r for r in registry}
        # custom big-employer sources (Amazon, etc.) that aren't standard ATS feeds
        try:
            from custom_sources import CUSTOM_FETCHERS
            for name, fn in CUSTOM_FETCHERS.items():
                futs[ex.submit(fn)] = {"company": name}
        except Exception:
            pass
        for fut in as_completed(futs):
            try:
                all_jobs.extend(fut.result() or [])
            except Exception:
                pass

    new_rows = []
    for j in all_jobs:
        url = j.get("url", "")
        if j["uid"] in seen or (url and url in seen):
            continue
        # pre-curated new-grad sources classify by role family only; others strict
        cls = classify_role(j["title"]) if j.get("prefiltered_newgrad") else classify(j["title"])
        if not cls:
            continue
        if not location_ok(j["location"]):
            continue
        seen.add(j["uid"])       # dedupe by source id
        if url:
            seen.add(url)        # and by URL, so two sources don't double-list
        category, level = cls
        new_rows.append({**j, "category": category, "level": level})

    # enrich comp for the (small number of) matched rows in parallel
    if new_rows:
        with ThreadPoolExecutor(max_workers=10) as ex:
            for j, comp in zip(new_rows, ex.map(enrich_comp, new_rows)):
                j["comp"] = comp or j.get("comp", "")

    now = datetime.now(ET).strftime("%Y-%m-%d %H:%M")
    if new_rows:
        wb = ensure_workbook()
        ws = wb["Jobs"]
        # newest on top: insert after header
        for j in sorted(new_rows, key=lambda x: x["company"]):
            ws.insert_rows(2)
            row = [now, j["company"], j["title"], j["category"], j["level"],
                   j["location"], j["comp"], j["posted"], j["url"], ""]
            for col, val in enumerate(row, 1):
                ws.cell(row=2, column=col, value=val)
            ws.cell(row=2, column=9).hyperlink = j["url"]
            ws.cell(row=2, column=9).font = Font(color="2563EB", underline="single")
        wb.save(XLSX_PATH)

    save_seen(seen)
    send_email_alert(new_rows)
    if new_rows and sys.platform == "darwin":
        try:
            import subprocess
            top = sorted(new_rows, key=lambda x: x["category"] != "APM")[0]
            body = f"{top['company']}: {top['title'][:60]}"
            if len(new_rows) > 1:
                body += f" (+{len(new_rows) - 1} more)"
            subprocess.run(["osascript", "-e",
                            f'display notification "{body}" with title '
                            f'"Job Tracker — {len(new_rows)} new" sound name "Glass"'],
                           timeout=10, check=False)
        except Exception:
            pass
    print(f"[{now}] scanned {len(all_jobs)} live postings across {len(registry)} companies -> "
          f"{len(new_rows)} NEW matching roles added. (sheet: {XLSX_PATH})")
    for j in sorted(new_rows, key=lambda x: (x['category'], x['company']))[:40]:
        comp = f"  [{j['comp']}]" if j['comp'] else ""
        print(f"    + {j['category']:<22} {j['company']:<14} {j['title'][:50]:<50} {j['location'][:24]}{comp}")


if __name__ == "__main__":
    main()
