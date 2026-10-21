#!/usr/bin/env python3
"""
Remove closed postings from data/jobs.xlsx AND the Google Sheet "Jobs" tab.

A row is only deleted on positive evidence that the posting is closed:
  1. ATS company (registry.json): that company's feed fetched successfully and
     returned postings, but this URL is no longer among them.
  2. Simplify list: the listing is present but marked inactive / hidden.
  3. Anything else (Amazon, Simplify rows that vanished, unknown source): the URL
     itself returns 404/410, or Greenhouse redirects to its "?error=true" page.
A failed fetch or a network error means "unknown" -> the row is kept. So a flaky
cycle can never wipe rows. Sheet rows with a Status (e.g. "Applied") are kept.

Usage: python3 prune_closed.py            # dry run: print what would go
       python3 prune_closed.py --apply    # actually delete
Run from run.sh at most every PRUNE_EVERY_HOURS (it refetches every feed).
"""
import json
import os
import re
import sys
import socket
from concurrent.futures import ThreadPoolExecutor
from urllib.request import Request, urlopen
from urllib.error import HTTPError

socket.setdefaulttimeout(60)  # same no-hang rule as the sheet scripts

import poller
import custom_sources

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX_PATH = os.path.join(HERE, "data", "jobs.xlsx")
KEY_PATH = os.path.join(HERE, "gcp-service-account.json")
CONFIG_PATH = os.path.join(HERE, "sheet_config.json")
STAMP_PATH = os.path.join(HERE, "data", "last_prune.txt")
PRUNE_EVERY_HOURS = 6
URL_COL, STATUS_COL = 9, 10  # 1-indexed, same layout in xlsx and Sheet


def log(msg):
    print(f"[prune_closed] {msg}")


def fetch_ats():
    """-> (live_urls, companies whose feed came back non-empty)"""
    registry = json.load(open(os.path.join(HERE, "registry.json")))["found"]
    live, ok = set(), set()

    def one(r):
        try:
            return r["company"], poller.FETCHERS[r["provider"]](r["company"], r["slug"]) or []
        except Exception:
            return r["company"], []

    with ThreadPoolExecutor(12) as ex:
        for company, jobs in ex.map(one, registry):
            if jobs:
                ok.add(company.strip().lower())
                live.update(j["url"].strip() for j in jobs if j.get("url"))
    return live, ok


def fetch_simplify():
    """-> {url: is_open} for every Simplify listing, open or not (None on failure)."""
    d = custom_sources._get(custom_sources.SIMPLIFY_URL)
    if not isinstance(d, list):
        return None
    return {(j.get("url") or "").strip(): bool(j.get("active") and j.get("is_visible"))
            for j in d if j.get("url")}


_ASHBY = {}


def check_url(url):
    """Ask the posting's own site. -> "open" | "closed" | "unknown" (any doubt)."""
    try:
        m = re.match(r"https://([^.]+)\.(wd\d+)\.myworkdayjobs\.com/(?:[a-z]{2}-[A-Z]{2}/)?"
                     r"([^/]+)/job/([^?#]+)", url)
        if m:  # Workday JSON API: 200 = live, 403/404/410 = gone
            tenant, wd, site, path = m.groups()
            return _status(f"https://{tenant}.{wd}.myworkdayjobs.com/wday/cxs/{tenant}/"
                           f"{site}/job/{path}", {"Accept": "application/json"})
        m = re.match(r"https://jobs\.ashbyhq\.com/([^/]+)/([0-9a-f-]{36})", url)
        if m:  # Ashby: is the id on the org's public board?
            org, jid = m.groups()
            if org not in _ASHBY:
                d = poller.fetch_json(f"https://api.ashbyhq.com/posting-api/job-board/{org}")
                _ASHBY[org] = {j.get("id") for j in d["jobs"]} if isinstance(d, dict) else None
            ids = _ASHBY[org]
            return "unknown" if ids is None else ("open" if jid in ids else "closed")
        m = re.match(r"https://(?:job-)?boards(?:\.eu)?\.greenhouse\.io/(?:embed/job_app\?for=)?([^/?]+)/jobs/(\d+)", url)
        if m:  # Greenhouse job API: 200 = live, 404 = gone
            return _status("https://boards-api.greenhouse.io/v1/boards/%s/jobs/%s" % m.groups())
        m = re.match(r"https://jobs\.smartrecruiters\.com/([^/]+)/(\d+)", url)
        if m:
            return _status("https://api.smartrecruiters.com/v1/companies/%s/postings/%s" % m.groups())
        if re.match(r"https://(jobs\.lever\.co|apply\.workable\.com|(job-)?boards(\.eu)?\.greenhouse\.io)/", url):
            return _status(url)
        return "closed" if _status(url) == "closed" else "unknown"
    except Exception:
        return "unknown"


def _status(url, headers=None):
    try:
        req = Request(url, headers={**poller.UA, **(headers or {})})
        with urlopen(req, timeout=20, context=poller.SSL_CTX) as r:
            final = r.geturl()
            # Greenhouse / Workable redirect closed jobs to an error page
            return "closed" if ("error=true" in final or "not_found=true" in final) else "open"
    except HTTPError as e:
        return "closed" if e.code in (403, 404, 410) and "myworkdayjobs" in url or e.code in (404, 410) \
            else "unknown"
    except Exception:
        return "unknown"


IDS = re.compile(r"\d{7,}")


def closed_urls(rows):
    """rows: [(company, url)] -> set of urls proven closed.

    Order of evidence: live in a feed -> keep. Otherwise ask the posting's own
    site (authoritative either way). Only if the site can't tell us do we fall
    back to "missing from a feed that loaded fine" or "Simplify marks inactive"
    (spot-checked ~95% right; the direct check catches the rest where possible)."""
    ats_live, ats_ok = fetch_ats()
    simplify = fetch_simplify()
    log(f"ATS feeds ok: {len(ats_ok)}; Simplify: {'ok' if simplify else 'FAILED'}")
    live_ids = {i for u in ats_live for i in IDS.findall(u)}
    suspects = []
    for company, url in rows:
        if url in ats_live or (simplify and simplify.get(url)):
            continue
        if company.strip().lower() in ats_ok and set(IDS.findall(url)) & live_ids:
            continue  # same job id, the site just changed its URL format
        suspects.append((company, url))
    with ThreadPoolExecutor(16) as ex:
        verdicts = list(ex.map(lambda cu: check_url(cu[1]), suspects))
    closed, counts = set(), {"open": 0, "closed": 0, "fallback": 0, "kept": 0}
    for (company, url), v in zip(suspects, verdicts):
        if v == "closed":
            closed.add(url); counts["closed"] += 1
        elif v == "open":
            counts["open"] += 1
        elif company.strip().lower() in ats_ok or (simplify is not None and url in simplify):
            closed.add(url); counts["fallback"] += 1
        else:
            counts["kept"] += 1
    log(f"checked {len(suspects)} suspects: site says closed {counts['closed']}, "
        f"site says open {counts['open']}, closed by feed evidence {counts['fallback']}, "
        f"kept as unknown {counts['kept']}")
    return closed


def main(apply):
    from openpyxl import load_workbook
    wb = load_workbook(XLSX_PATH)
    ws = wb["Jobs"]
    rows = []
    for r in range(2, ws.max_row + 1):
        url = str(ws.cell(r, URL_COL).value or "").strip()
        if url:
            rows.append((str(ws.cell(r, 2).value or ""), url))
    closed = closed_urls(rows)

    # Sheet first: keep anything you've put a Status on.
    sheet_ws, keep = None, set()
    if os.path.exists(KEY_PATH) and os.path.exists(CONFIG_PATH):
        import gspread
        from google.oauth2.service_account import Credentials
        creds = Credentials.from_service_account_file(
            KEY_PATH, scopes=["https://www.googleapis.com/auth/spreadsheets"])
        sheet_id = json.load(open(CONFIG_PATH))["sheet_id"]
        sheet_ws = gspread.authorize(creds).open_by_key(sheet_id).worksheet("Jobs")
        values = sheet_ws.get_all_values()
        keep = {r[URL_COL - 1].strip() for r in values[1:]
                if len(r) >= STATUS_COL and r[STATUS_COL - 1].strip()}
    closed -= keep

    for company, url in rows:
        if url in closed:
            log(f"closed: {company} | {url}")
    log(f"{len(closed)} of {len(rows)} postings closed"
        + (f" ({len(keep)} kept because they have a Status)" if keep else ""))
    if not apply or not closed:
        return

    if sheet_ws is not None:
        dead_rows = [i + 1 for i, r in enumerate(values) if i > 0
                     and len(r) >= URL_COL and r[URL_COL - 1].strip() in closed]
        # delete bottom-up so row numbers stay valid; one API call
        reqs = [{"deleteDimension": {"range": {"sheetId": sheet_ws.id, "dimension": "ROWS",
                                               "startIndex": n - 1, "endIndex": n}}}
                for n in sorted(dead_rows, reverse=True)]
        if reqs:
            sheet_ws.spreadsheet.batch_update({"requests": reqs})
        log(f"deleted {len(dead_rows)} rows from the Sheet")

    for r in range(ws.max_row, 1, -1):
        if str(ws.cell(r, URL_COL).value or "").strip() in closed:
            ws.delete_rows(r)
    wb.save(XLSX_PATH)
    log(f"deleted {len(closed)} rows from jobs.xlsx")


def due():
    import time
    try:
        return time.time() - os.path.getmtime(STAMP_PATH) > PRUNE_EVERY_HOURS * 3600
    except OSError:
        return True


if __name__ == "__main__":
    apply = "--apply" in sys.argv
    if "--if-due" in sys.argv and not due():
        sys.exit(0)
    try:
        main(apply)
        if apply:
            open(STAMP_PATH, "w").write("ok\n")
    except Exception as e:
        log(f"ERROR: {e}")
        sys.exit(0)  # never break the poller pipeline
